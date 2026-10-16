import json
import os
import resource
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

from omisa_parser.core import Limits, ParseError, parse_document
from omisa_parser.protocol import ProtocolError, validate_response_stream

RID = "a" * 32
FIX = Path(__file__).resolve().parents[1] / "fixtures"


def _raw(rows):
    return [(json.dumps(row, separators=(",", ":")) + "\n").encode() for row in rows]


def _end(**changes):
    row = {"schema":"omisa.document/1","type":"end","request_id":RID,"status":"ok","counts":{"units_emitted":0,"records_emitted":1,"text_chars":0,"tables":0,"cells":0},"truncated":False,"limit_reason":None,"warnings":[]}
    row.update(changes)
    return row


def test_protocol_validates_schema_state_counts_references_and_warnings():
    document={"schema":"omisa.document/1","type":"document","request_id":RID,"document":{"detected_format":"txt","media_type":"text/plain","size_bytes":1,"title":None,"language":None},"parser":{"bundle_version":"x","registry_version":1},"structure":{"unit":"page","unit_count":1,"pagination":"logical"},"warnings":[{"code":"active_content_present","count":1}]}
    unit={"schema":"omisa.document/1","type":"unit","id":"u1","unit_kind":"page","index":1,"label":"1","parent_id":None}
    text={"schema":"omisa.document/1","type":"text","id":"b1","unit_id":"u1","order":1,"text":"x","source":"native","truncated":False}
    end=_end(counts={"units_emitted":1,"records_emitted":4,"text_chars":1,"tables":0,"cells":0},warnings=document["warnings"])
    assert len(list(validate_response_stream(_raw([document,unit,text,end]), request_id=RID))) == 4
    bad=[]
    bad.append([dict(document, surprise=True), end])
    bad.append([document, dict(text, unit_id="u9"), end])
    bad.append([document,unit,text,dict(end, counts=dict(end["counts"], text_chars=2))])
    bad.append([document,unit,text,dict(end, warnings=[])])
    bad.append([document,unit,dict(text, order=1),dict(text,id="b2",order=1),end])
    for rows in bad:
        with pytest.raises(ProtocolError): list(validate_response_stream(_raw(rows), request_id=RID))


def test_json_and_rtf_are_rejected_before_large_full_materialization(tmp_path):
    limits=Limits(max_input_bytes=20*1024*1024, max_json_bytes=1024*1024, max_rtf_bytes=1024*1024)
    for prefix in (b'{"x":"', b'{\\rtf1 '):
        p=tmp_path/str(len(prefix));p.write_bytes(prefix+b'x'*(2*1024*1024)+b'"}' if prefix.startswith(b'{"') else prefix+b'x'*(2*1024*1024)+b'}')
        with pytest.raises(ParseError) as exc: parse_document(p,RID,limits=limits)
        assert exc.value.code == 'resource_limit'


def test_requested_range_is_applied_before_pdf_sheet_slide_extraction(monkeypatch):
    import pymupdf
    seen=[]
    real=pymupdf.Page.get_text
    def guard(self,*a,**kw): seen.append(self.number+1); return real(self,*a,**kw)
    monkeypatch.setattr(pymupdf.Page,'get_text',guard)
    parse_document(FIX/'sample.pdf',RID,{'range':{'unit':'page','start':1,'end':1}})
    assert seen == [1]


def test_global_text_budget_includes_notes_cells_and_formulas(tmp_path):
    from openpyxl import Workbook
    p=tmp_path/'x.xlsx'; wb=Workbook(); ws=wb.active; ws['A1']='abcd'; ws['A2']='=12345'; wb.save(p)
    with pytest.raises(ParseError) as exc: parse_document(p,RID,limits=Limits(max_text_chars=8))
    assert exc.value.code == 'resource_limit'


def test_formula_cache_unavailable_and_external_dde_are_classified(tmp_path):
    from openpyxl import Workbook
    p=tmp_path/'x.xlsx'; wb=Workbook(); ws=wb.active
    ws['A1']='=1+1'; ws['A2']="='[evil.xlsx]S'!A1"; ws['A3']='=cmd|\' /C calc\'!A0'; wb.save(p)
    records=parse_document(p,RID)
    warnings={w['code']:w['count'] for w in records[0]['warnings']}
    assert warnings == {'formula_cache_unavailable':3,'external_formula_ignored':1,'dde_formula_ignored':1}
    cells=next(r for r in records if r['type']=='table')['cells']
    assert all(c['formula'] and c['value']['kind']=='blank' for c in cells)


def test_odf_repeated_rows_columns_checked_product_and_expansion(tmp_path):
    p=tmp_path/'x.ods'
    content='''<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:spreadsheet><table:table table:name="S"><table:table-row table:number-rows-repeated="3"><table:table-cell table:number-columns-repeated="4"><text:p>x</text:p></table:table-cell></table:table-row></table:table></office:spreadsheet></office:body></office:document-content>'''
    manifest='''<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/><manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/></manifest:manifest>'''
    with zipfile.ZipFile(p,'w') as z:
        z.writestr('mimetype','application/vnd.oasis.opendocument.spreadsheet',compress_type=zipfile.ZIP_STORED)
        z.writestr('META-INF/manifest.xml',manifest)
        z.writestr('content.xml',content)
    records=parse_document(p,RID,limits=Limits(max_cells=12))
    table=next(r for r in records if r['type']=='table')
    assert (table['row_count'],table['column_count'],len(table['cells'])) == (3,4,12)
    with pytest.raises(ParseError): parse_document(p,RID,limits=Limits(max_cells=11))


@pytest.mark.parametrize('kind', ['xlsx','pptx'])
def test_macro_spreadsheet_and_presentation_variants_warn(kind,tmp_path):
    source=FIX/f'sample.{kind}'; target=tmp_path/f'x.{kind}'
    replacements={
      'xlsx':(b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml',b'application/vnd.ms-excel.sheet.macroEnabled.main+xml','xl/vbaProject.bin'),
      'pptx':(b'application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml',b'application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml','ppt/vbaProject.bin')}
    old,new,member=replacements[kind]
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(target,'w') as dst:
        for info in src.infolist():
            data=src.read(info)
            if info.filename=='[Content_Types].xml': data=data.replace(old,new)
            dst.writestr(info,data)
        dst.writestr(member,b'macro')
    records=parse_document(target,RID)
    assert any(w['code']=='active_content_present' for w in records[0]['warnings'])
