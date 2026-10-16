import io, json, zipfile
from pathlib import Path
import pytest
from omisa_parser.core import Limits, ParseError, parse_document

RID='a'*32
FIX=Path(__file__).parents[1]/'fixtures'

def assert_code(path, code, **kw):
    with pytest.raises(ParseError) as e: parse_document(path,RID,**kw)
    assert e.value.code==code

def test_refuses_generic_archive(tmp_path):
    p=tmp_path/'renamed.docx'
    with zipfile.ZipFile(p,'w') as z:z.writestr('payload.txt','no')
    assert_code(p,'unsupported_format')

def test_refuses_nested_archive(tmp_path):
    p=tmp_path/'x.docx'
    with zipfile.ZipFile(p,'w') as z:
        z.writestr('[Content_Types].xml',b'wordprocessingml.document')
        z.writestr('word/embedded.zip',b'PK\x03\x04')
    assert_code(p,'archive_policy_violation')

def test_refuses_zip_traversal(tmp_path):
    p=tmp_path/'x.docx'
    with zipfile.ZipFile(p,'w') as z:z.writestr('../evil','x')
    assert_code(p,'archive_policy_violation')

def test_refuses_high_ratio_zip(tmp_path):
    p=tmp_path/'x.docx'
    with zipfile.ZipFile(p,'w',zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml',b'wordprocessingml.document')
        z.writestr('word/document.xml',b'A'*200000)
    assert_code(p,'archive_policy_violation')

def test_refuses_xml_dtd_and_entities(tmp_path):
    p=tmp_path/'safe.txt';p.write_text('<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x>&e;</x>')
    assert_code(p,'active_content_policy')

def test_refuses_malformed_pdf(tmp_path):
    p=tmp_path/'x';p.write_bytes(b'%PDF-garbage')
    assert_code(p,'malformed_document')


def test_utf8_detection_window_may_end_mid_character(tmp_path):
    p=tmp_path/'x.md';p.write_bytes(b'# Heading\n'+b'a'*8180+'Ж'.encode()+b'\nrest')
    assert parse_document(p,RID)[0]['document']['detected_format']=='md'


def test_pdf_additional_actions_are_semantic_not_raw_substrings(tmp_path):
    source=FIX/'sample.pdf'
    data=source.read_bytes()
    safe=tmp_path/'safe.pdf';safe.write_bytes(data.replace(b'%%EOF',b'<</FontName/AAAAAA+Roboto-Regular>>\n%%EOF'))
    assert parse_document(safe,'1'*32)[-1]['type']=='end'
    import pymupdf
    doc=pymupdf.open();page=doc.new_page();action=doc.get_new_xref()
    doc.update_object(action,"<</S/JavaScript/JS(app.alert)>>")
    doc.xref_set_key(page.xref,"AA",f"<</O {action} 0 R>>")
    active=tmp_path/'active.pdf';active.write_bytes(doc.tobytes());doc.close()
    assert_code(active,'active_content_policy')


def test_pdf_slash_aa_inside_compressed_content_stream_is_ignored(tmp_path):
    import pymupdf
    doc=pymupdf.open();page=doc.new_page();page.insert_text((72,72),"/AA /JS /JavaScript")
    safe=tmp_path/'safe.pdf';safe.write_bytes(doc.tobytes());doc.close()
    assert parse_document(safe,'1'*32,{'ocr':'off'})[-1]['type']=='end'

def test_input_cap(tmp_path):
    p=tmp_path/'x';p.write_bytes(b'x'*11)
    assert_code(p,'input_too_large',limits=Limits(max_input_bytes=10))

def test_output_is_bounded_and_terminal(tmp_path):
    p=tmp_path/'x';p.write_text('x'*10000)
    records=parse_document(p,RID,limits=Limits(max_output_bytes=1000))
    raw=('\n'.join(json.dumps(x,separators=(',',':')) for x in records)+'\n').encode()
    assert len(raw)<=1000
    assert records[-1]['type']=='end' and records[-1]['truncated']

def test_controls_are_sanitized(tmp_path):
    p=tmp_path/'x';p.write_text('safe\x00\u202eevil')
    records=parse_document(p,RID)
    text=next(r['text'] for r in records if r['type']=='text')
    assert '\x00' not in text and '\u202e' not in text

def test_range_and_record_cap(tmp_path):
    p=tmp_path/'x';p.write_text('ok')
    records=parse_document(p,RID,options={'max_records':2})
    assert len(records)<=3


def test_requested_max_records_above_default_is_honored(tmp_path):
    from docx import Document
    p=tmp_path/'x.docx';d=Document()
    for i in range(250):d.add_paragraph(f'Paragraph {i}')
    d.save(p)
    records=parse_document(p,RID,options={'max_records':500})
    assert len([row for row in records if row['type']=='text'])==250
    assert records[-1]['limit_reason'] is None


def test_internal_record_offset_pages_large_docx_without_gaps(tmp_path):
    from docx import Document
    p=tmp_path/'large.docx';d=Document()
    for i in range(620):d.add_paragraph(f'Paragraph {i:04d}')
    d.save(p)
    first=parse_document(p,RID,options={'max_records':200,'record_offset':0})
    first_text=[row for row in first if row['type']=='text']
    assert first[-1]['truncated'] is True and first[-1]['limit_reason']=='records'
    second=parse_document(p,'b'*32,options={'max_records':200,'record_offset':len(first_text)})
    second_text=[row for row in second if row['type']=='text']
    assert second_text[0]['text']==f'Paragraph {len(first_text):04d}'
    seen=[row['text'] for row in first_text+second_text]
    assert len(seen)==len(set(seen))


def test_large_csv_table_is_split_into_resumable_row_ranges(tmp_path):
    p=tmp_path/'large.csv'
    p.write_text('name,value\n'+''.join(f'row{i:04d},{"x"*100}\n' for i in range(1500)))
    offset=0;chunks=[]
    for n in range(20):
        rows=parse_document(p,f'{n:032x}',options={'max_records':50,'record_offset':offset})
        tables=[row for row in rows if row['type']=='table']
        chunks.extend(tables);offset+=len(tables)
        if not rows[-1]['truncated']:break
    assert rows[-1]['truncated'] is False
    assert len(chunks)>1
    assert [(row['row_start'],row['row_end']) for row in chunks][0][0]==1
    assert chunks[-1]['row_end']==1501
    assert all(a['row_end']+1==b['row_start'] for a,b in zip(chunks,chunks[1:]))
    assert sum(row['row_count'] for row in chunks)==1501
    assert all(len((json.dumps(row,ensure_ascii=False,separators=(',',':'))+'\n').encode())<64*1024 for row in chunks)

def test_xlsx_cell_cap_is_enforced_before_output(tmp_path):
    from openpyxl import Workbook
    p=tmp_path/'many.xlsx'; w=Workbook(); ws=w.active
    ws.append([1,2]); ws.append([3,4]); w.save(p)
    assert_code(p,'resource_limit',limits=Limits(max_cells=3))
